"""Evidence-preserving PDF/DOCX ingestion tests."""

from __future__ import annotations

from pathlib import Path

import fitz
import pytest
from docx import Document as DocxDocument

from core.document_loader import DocumentError, DocumentLoader
from tests.helpers import sample_path

SAMPLE_DOCX = "demo_contract.docx"
SAMPLE_PDF = "demo_contract.pdf"
SCANNED_PDF = "scanned_no_text.pdf"


def test_docx_offsets_match_every_segment(pipeline):
    document = pipeline.loader.load(str(sample_path(SAMPLE_DOCX)))
    assert document.file_type == "docx"
    assert document.stats.extraction_ok
    assert document.stats.num_paragraphs >= 10
    assert document.raw_text

    segments = pipeline.splitter.split(document)
    assert segments
    for segment in segments:
        span = document.span_by_id(segment.source_id)
        assert span is not None
        assert 0 <= segment.span_start < segment.span_end <= len(span.text)
        assert span.text[segment.span_start : segment.span_end].strip() == segment.text
        location = document.location_by_id(segment.source_id)
        assert location is not None
        assert location.paragraph_index == span.paragraph_index


def test_docx_table_rows_are_preserved(tmp_path: Path):
    path = tmp_path / "table_contract.docx"
    doc = DocxDocument()
    doc.add_paragraph("1. Payment. The Customer shall pay each invoice within 30 days.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Term"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Governing law"
    table.cell(1, 1).text = "State of Delaware"
    doc.save(path)

    loaded = DocumentLoader().load(path)
    table_spans = [span for span in loaded.spans if span.source_type == "docx_table_row"]
    assert table_spans
    assert any("Governing law" in span.text and "Delaware" in span.text for span in table_spans)
    assert loaded.stats.num_tables == 1
    for span in table_spans:
        assert loaded.location_by_id(span.source_id).table_index is not None
        assert loaded.location_by_id(span.source_id).row_index is not None


def test_pdf_pages_and_locations_are_preserved():
    document = DocumentLoader().load(sample_path(SAMPLE_PDF))
    assert document.file_type == "pdf"
    assert document.stats.num_pages and document.stats.num_pages >= 2
    assert document.stats.extraction_ok
    assert document.spans

    page_numbers = {span.page_number for span in document.spans}
    assert page_numbers == set(range(1, document.stats.num_pages + 1))
    for span in document.spans:
        location = document.location_by_id(span.source_id)
        assert location is not None
        assert location.page_number == span.page_number
        assert location.source_type == "pdf_block"
        assert span.source_id in document.source_map


def test_scanned_pdf_is_detected_not_silently_empty():
    document = DocumentLoader().load(sample_path(SCANNED_PDF))
    assert document.file_type == "pdf"
    assert document.stats.needs_ocr is True
    assert document.stats.extraction_ok is False
    assert document.stats.extraction_coverage == 0.0
    assert any("OCR" in warning or "image" in warning.lower() for warning in document.stats.warnings)


def test_text_ingestion_and_empty_text():
    loader = DocumentLoader()
    document = loader.load_text(
        "1. Payment. Invoices are due within 30 days.\n\n"
        "2. Governing law. This Agreement is governed by the laws of Delaware."
    )
    assert document.file_type == "text"
    assert len(document.spans) == 2
    assert document.spans[0].paragraph_index == 1
    assert document.stats.extraction_ok

    empty = loader.load_text("   \n\n  ")
    assert empty.spans == []
    assert empty.stats.extraction_ok is False
    assert empty.stats.extraction_coverage == 0.0
    assert empty.stats.warnings


def test_empty_oversized_unsupported_and_bad_signatures(tmp_path: Path):
    loader = DocumentLoader(max_file_size_mb=1)

    empty = tmp_path / "empty.pdf"
    empty.write_bytes(b"")
    with pytest.raises(DocumentError, match="empty"):
        loader.load(empty)

    oversized = tmp_path / "large.pdf"
    oversized.write_bytes(b"%PDF-1.4\n" + b"x" * (1024 * 1024 + 1))
    with pytest.raises(DocumentError, match="limit"):
        loader.load(oversized)

    legacy = tmp_path / "contract.doc"
    legacy.write_bytes(b"\xd0\xcf\x11\xe0")
    with pytest.raises(DocumentError, match="Legacy"):
        loader.load(legacy)

    unsupported = tmp_path / "contract.rtf"
    unsupported.write_text("{\\rtf1 contract}", encoding="utf-8")
    with pytest.raises(DocumentError, match="Unsupported"):
        loader.load(unsupported)

    bad_pdf = tmp_path / "fake.pdf"
    bad_pdf.write_bytes(b"this is not a pdf")
    with pytest.raises(DocumentError, match="PDF"):
        loader.load(bad_pdf)

    bad_docx = tmp_path / "fake.docx"
    bad_docx.write_bytes(b"PK\x03\x04not really a zip package")
    with pytest.raises(DocumentError):
        loader.load(bad_docx)


def test_pdf_page_truncation_is_explicit(tmp_path: Path):
    path = tmp_path / "two_page.pdf"
    pdf = fitz.open()
    for page_number in (1, 2):
        page = pdf.new_page()
        page.insert_text(
            (72, 72),
            f"Page {page_number}. Payment shall be made within 30 days of invoice.",
        )
    pdf.save(path)
    pdf.close()

    document = DocumentLoader(max_pdf_pages=1).load(path)
    assert document.stats.truncated is True
    assert document.stats.num_pages == 2
    assert any("first 1 of 2" in warning for warning in document.stats.warnings)
    assert all(span.page_number == 1 for span in document.spans)


def test_empty_analysis_produces_not_verified_coverage(pipeline):
    result = pipeline.analyze_text("", "empty.txt")
    assert result.document.spans == []
    assert result.summary.categories_detected == 0
    assert result.summary.not_verified == 4
    assert all(finding.status.value == "NOT_VERIFIED" for finding in result.findings)
    assert all(not finding.has_evidence for finding in result.findings)
