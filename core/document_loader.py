"""Evidence-preserving PDF and DOCX ingestion.

PDF pages are extracted with PyMuPDF and split into text blocks; DOCX content is
read in body order including tables. Every extracted unit receives a stable
source identifier plus page or paragraph metadata. The raw extracted text is
never rewritten for matching.
"""

from __future__ import annotations

import os
import re
from pathlib import Path
from typing import Iterable, Optional

import fitz  # PyMuPDF
from docx import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph

from core.text_utils import normalize_whitespace, sha256_bytes
from models.schemas import (
    Document,
    DocumentSpan,
    DocumentStats,
    SourceLocation,
)


class DocumentError(ValueError):
    """Raised when a document cannot be safely or reliably extracted."""


class DocumentLoader:
    SUPPORTED_EXTENSIONS = {"pdf", "docx"}
    INTERNAL_EXTENSIONS = {"md", "txt"}

    def __init__(self, max_file_size_mb: int = 10, max_pdf_pages: int = 300) -> None:
        if max_file_size_mb <= 0:
            raise ValueError("max_file_size_mb must be positive")
        self.max_file_size_bytes = max_file_size_mb * 1024 * 1024
        self.max_pdf_pages = max_pdf_pages

    # ------------------------------------------------------------------ public
    def load(self, file_path: str | Path, filename: Optional[str] = None) -> Document:
        path = Path(file_path)
        display_name = filename or path.name
        ext = self._extension(display_name)
        if not path.exists():
            raise DocumentError("The uploaded file could not be found after saving.")
        if not path.is_file():
            raise DocumentError("The uploaded path is not a file.")
        size = path.stat().st_size
        if size == 0:
            raise DocumentError("The uploaded file is empty.")
        if size > self.max_file_size_bytes:
            raise DocumentError(
                f"File is {size / (1024 * 1024):.1f} MB; the limit is "
                f"{self.max_file_size_bytes / (1024 * 1024):.0f} MB."
            )

        data = path.read_bytes()
        digest = sha256_bytes(data)

        if ext == "pdf":
            self._validate_pdf_signature(data)
            return self._load_pdf(path, display_name, digest)
        if ext == "docx":
            self._validate_docx_signature(data)
            return self._load_docx(path, display_name, digest)
        if ext in self.INTERNAL_EXTENSIONS:
            text = data.decode("utf-8", errors="replace")
            return self.load_text(text, display_name, sha256=digest)
        if ext == "doc":
            raise DocumentError(
                "Legacy .doc files are not supported. Save the contract as .docx first."
            )
        raise DocumentError(
            f"Unsupported file type '.{ext}'. Upload a PDF or DOCX contract."
        )

    def load_text(self, text: str, filename: str = "pasted-contract.md", sha256: Optional[str] = None) -> Document:
        """Internal text ingestion used by samples and tests."""
        paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text or "") if p.strip()]
        spans: list[DocumentSpan] = []
        for index, paragraph in enumerate(paragraphs, start=1):
            spans.append(
                DocumentSpan(
                    source_id=f"text-p{index}",
                    text=paragraph,
                    paragraph_index=index,
                    source_type="text",
                    char_start=0,
                    char_end=len(paragraph),
                )
            )
        raw_text = "\n\n".join(span.text for span in spans)
        warnings: list[str] = []
        if not spans:
            warnings.append("No readable text was found in the document.")
        stats = DocumentStats(
            filename=filename,
            file_type="text",
            num_paragraphs=len(spans),
            num_chars=len(raw_text),
            num_words=len(raw_text.split()),
            pages_with_text=None,
            extraction_coverage=1.0 if spans else 0.0,
            extraction_ok=bool(spans),
            warnings=warnings,
        )
        return Document(
            filename=filename,
            file_type="text",
            raw_text=raw_text,
            spans=spans,
            stats=stats,
            source_map={
                span.source_id: SourceLocation(
                    source_id=span.source_id,
                    paragraph_index=span.paragraph_index,
                    source_type="text",
                )
                for span in spans
            },
            sha256=sha256,
        )

    # ------------------------------------------------------------------ PDF
    def _load_pdf(self, path: Path, filename: str, digest: str) -> Document:
        try:
            pdf = fitz.open(path)
        except Exception as exc:  # pragma: no cover - depends on malformed input
            raise DocumentError(f"Could not open the PDF: {exc}") from exc

        try:
            if pdf.needs_pass:
                raise DocumentError(
                    "The PDF is password-protected. Remove the password and try again."
                )
            total_pages = pdf.page_count
            if total_pages == 0:
                raise DocumentError("The PDF has no pages.")

            truncated = total_pages > self.max_pdf_pages
            pages_to_process = min(total_pages, self.max_pdf_pages)
            spans: list[DocumentSpan] = []
            raw_parts: list[str] = []
            warnings: list[str] = []
            pages_with_text = 0
            blank_pages: list[int] = []
            block_sequence = 0

            for page_index in range(pages_to_process):
                page = pdf[page_index]
                page_number = page_index + 1
                blocks = page.get_text("blocks", sort=True)
                page_has_text = False
                block_number = 0
                for block in blocks:
                    # PyMuPDF text blocks: (x0, y0, x1, y1, text, block_no, block_type)
                    if len(block) < 7 or block[6] != 0:
                        continue
                    text = str(block[4]).strip()
                    if not text:
                        continue
                    block_number += 1
                    block_sequence += 1
                    page_has_text = True
                    spans.append(
                        DocumentSpan(
                            source_id=f"pdf-p{page_number}-b{block_number}",
                            text=text,
                            page_number=page_number,
                            block_index=block_number,
                            source_type="pdf_block",
                            char_start=0,
                            char_end=len(text),
                        )
                    )
                    raw_parts.append(text)
                if page_has_text:
                    pages_with_text += 1
                elif len(blank_pages) < 10:
                    blank_pages.append(page_number)

            raw_text = "\n\n".join(raw_parts)
            coverage = pages_with_text / pages_to_process if pages_to_process else 0.0
            needs_ocr = pages_with_text == 0 or coverage < 0.5
            extraction_ok = bool(raw_text.strip()) and len(raw_text.strip()) >= 40

            if truncated:
                warnings.append(
                    f"Only the first {pages_to_process} of {total_pages} PDF pages were "
                    "processed. Findings do not cover the remaining pages."
                )
            if blank_pages:
                shown = ", ".join(str(p) for p in blank_pages[:10])
                suffix = "" if len(blank_pages) <= 10 else ", …"
                warnings.append(
                    f"Pages with no extractable text (possibly scanned): {shown}{suffix}."
                )
            if needs_ocr:
                warnings.append(
                    "Low text coverage detected. OCR is not included in this MVP; "
                    "clauses on image-only pages cannot be verified."
                )
            if not extraction_ok:
                warnings.append(
                    "Very little readable text was extracted. Treat all categories as "
                    "NOT_VERIFIED rather than absent."
                )

            stats = DocumentStats(
                filename=filename,
                file_type="pdf",
                num_pages=total_pages,
                num_blocks=block_sequence,
                num_chars=len(raw_text),
                num_words=len(raw_text.split()),
                pages_with_text=pages_with_text,
                extraction_coverage=round(coverage, 4),
                extraction_ok=extraction_ok,
                truncated=truncated,
                warnings=warnings,
                needs_ocr=needs_ocr,
            )
            source_map = {
                span.source_id: SourceLocation(
                    source_id=span.source_id,
                    page_number=span.page_number,
                    block_index=span.block_index,
                    source_type="pdf_block",
                )
                for span in spans
            }
            return Document(
                filename=filename,
                file_type="pdf",
                raw_text=raw_text,
                spans=spans,
                stats=stats,
                source_map=source_map,
                sha256=digest,
            )
        finally:
            pdf.close()

    # ------------------------------------------------------------------ DOCX
    def _load_docx(self, path: Path, filename: str, digest: str) -> Document:
        try:
            document = DocxDocument(str(path))
        except Exception as exc:
            raise DocumentError(
                "Could not open the DOCX. The file may be corrupt or not a Word document."
            ) from exc

        spans: list[DocumentSpan] = []
        raw_parts: list[str] = []
        warnings: list[str] = []
        paragraph_number = 0
        table_number = 0
        row_count = 0
        section_id: Optional[str] = None

        for child in document.element.body.iterchildren():
            tag = child.tag.rsplit("}", 1)[-1]
            if tag == "p":
                paragraph = Paragraph(child, document)
                paragraph_number += 1
                text = paragraph.text.strip()
                style_name = ""
                try:
                    style_name = paragraph.style.name or ""
                except Exception:
                    style_name = ""
                if style_name.casefold().startswith("heading") and text:
                    section_id = text[:120]
                if not text:
                    continue
                spans.append(
                    DocumentSpan(
                        source_id=f"docx-para-{paragraph_number}",
                        text=text,
                        paragraph_index=paragraph_number,
                        section_id=section_id,
                        source_type="docx_paragraph",
                        char_start=0,
                        char_end=len(text),
                    )
                )
                raw_parts.append(text)
            elif tag == "tbl":
                table = Table(child, document)
                table_number += 1
                for row_number, row in enumerate(table.rows, start=1):
                    row_count += 1
                    cell_texts = [
                        normalize_whitespace(cell.text)
                        for cell in row.cells
                        if normalize_whitespace(cell.text)
                    ]
                    # Merged cells repeat text; remove consecutive duplicates.
                    deduped: list[str] = []
                    for cell_text in cell_texts:
                        if not deduped or deduped[-1] != cell_text:
                            deduped.append(cell_text)
                    text = " | ".join(deduped).strip()
                    if not text:
                        continue
                    spans.append(
                        DocumentSpan(
                            source_id=f"docx-table-{table_number}-row-{row_number}",
                            text=text,
                            table_index=table_number,
                            row_index=row_number,
                            section_id=section_id,
                            source_type="docx_table_row",
                            char_start=0,
                            char_end=len(text),
                        )
                    )
                    raw_parts.append(text)

        raw_text = "\n\n".join(raw_parts)
        paragraph_span_count = sum(1 for span in spans if span.paragraph_index is not None)
        table_span_count = sum(1 for span in spans if span.table_index is not None)
        extraction_ok = bool(raw_text.strip())
        coverage = 1.0 if extraction_ok else 0.0

        if not extraction_ok:
            warnings.append(
                "No readable paragraphs or table rows were found. The DOCX may be empty "
                "or contain only images."
            )
        elif len(raw_text.strip()) < 40:
            warnings.append("Very little text was extracted; manual review is required.")
        if table_span_count:
            warnings.append(
                f"{table_span_count} table row(s) extracted. Table cells are attributed "
                "by table and row, not page number."
            )
        if len(raw_text) > 300_000:
            warnings.append(
                "Large document: extraction is complete, but retrieval may take longer."
            )

        stats = DocumentStats(
            filename=filename,
            file_type="docx",
            num_paragraphs=paragraph_span_count,
            num_tables=table_number,
            num_blocks=len(spans),
            num_chars=len(raw_text),
            num_words=len(raw_text.split()),
            pages_with_text=None,
            extraction_coverage=coverage,
            extraction_ok=extraction_ok,
            warnings=warnings,
            needs_ocr=False,
        )
        source_map = {
            span.source_id: SourceLocation(
                source_id=span.source_id,
                paragraph_index=span.paragraph_index,
                table_index=span.table_index,
                row_index=span.row_index,
                source_type=span.source_type,
            )
            for span in spans
        }
        return Document(
            filename=filename,
            file_type="docx",
            raw_text=raw_text,
            spans=spans,
            stats=stats,
            source_map=source_map,
            sha256=digest,
        )

    # --------------------------------------------------------------- helpers
    @staticmethod
    def _extension(filename: str) -> str:
        name = (filename or "").strip().lower()
        if "." not in name:
            return ""
        return name.rsplit(".", 1)[-1]

    @staticmethod
    def _validate_pdf_signature(data: bytes) -> None:
        if b"%PDF-" not in data[:1024]:
            raise DocumentError("The file does not look like a valid PDF (missing %PDF header).")

    @staticmethod
    def _validate_docx_signature(data: bytes) -> None:
        if not data.startswith(b"PK\x03\x04"):
            raise DocumentError(
                "The file does not look like a valid DOCX (missing ZIP header)."
            )
